"""Production tests for the Q1 CLI and exports: files, read-back, mapping, template untouched."""
import csv
import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from microgrid import q1
from microgrid.export_q1 import LEDGER_COLUMNS, MAPPING_COLUMNS, read_schedule_csv

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "data/templates/result1.xlsx"


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


@pytest.fixture(scope="module")
def run_dir(tmp_path_factory):
    out = tmp_path_factory.mktemp("q1run")
    before = sha256(TEMPLATE)
    summary = q1.run(ROOT / "configs/q1.json", out, argv=["--config", "configs/q1.json", "--output", str(out)])
    assert sha256(TEMPLATE) == before, "template must never be modified"
    return out, summary


def test_outputs_exist(run_dir):
    out, _ = run_dir
    for name in ["schedule.csv", "summary.json", "result1.xlsx", "template_mapping.csv",
                 "tables_q1.md", "run_metadata.json"]:
        assert (out / name).is_file(), name


def test_schedule_csv_roundtrip(run_dir):
    out, summary = run_dir
    rows = read_schedule_csv(out / "schedule.csv")
    assert len(rows) == 144
    assert rows[0]["start"] == "00:00" and rows[0]["end"] == "00:10"
    assert rows[-1]["start"] == "23:50" and rows[-1]["end"] == "24:00"
    assert rows[60]["start"] == "10:00" and rows[60]["slot"] == 61
    total = sum(r["cost_yuan"] for r in rows)
    assert total == pytest.approx(summary["objective_yuan"], abs=1e-9)
    for r in rows:
        assert r["cost_yuan"] == pytest.approx(r["price_yuan_per_kwh"] * r["grid_kwh"], abs=1e-9)
    for a, b in zip(rows, rows[1:]):
        assert a["soc_end_kwh"] == pytest.approx(b["soc_start_kwh"], abs=1e-9)
    with (out / "schedule.csv").open(encoding="utf-8") as stream:
        assert next(csv.reader(stream)) == LEDGER_COLUMNS


def test_summary_contract(run_dir):
    out, summary = run_dir
    on_disk = json.loads((out / "summary.json").read_text(encoding="utf-8"))
    assert on_disk["schema_version"] == 1
    for key in ["objective_yuan", "total_grid_kwh", "total_charge_kwh", "total_discharge_kwh",
                "total_spill_kwh", "soc_initial_kwh", "soc_final_kwh", "input_sha256",
                "template_sha256", "solver", "source_git_commit", "source_dirty"]:
        assert key in on_disk, key
    assert on_disk["template_sha256"] == sha256(TEMPLATE)
    assert on_disk["solver"]["status"] == "optimal"
    assert 0 <= on_disk["solver"]["mip_gap"] <= 1e-8
    assert on_disk["soc_initial_kwh"] == 6000 and on_disk["soc_final_kwh"] == pytest.approx(6000, abs=1e-6)
    assert isinstance(on_disk["source_dirty"], bool)
    assert [s["slot"] for s in on_disk["selected_intervals"]] == [61, 73, 85, 97, 109, 121]
    assert on_disk["selected_intervals"][0]["interval"] == "10:00-10:10"
    meta = json.loads((out / "run_metadata.json").read_text(encoding="utf-8"))
    assert {"python", "numpy", "scipy", "openpyxl"} <= set(meta["versions"])
    assert meta["parameters"]["max_charge_kwh_per_slot"] == pytest.approx(5000 / 6)


def test_workbook_readback(run_dir):
    out, _ = run_dir
    rows = read_schedule_csv(out / "schedule.csv")
    wb = load_workbook(out / "result1.xlsx")
    try:
        assert wb.sheetnames == ["计划购电量", "充放电量"]
        plan, store = wb["计划购电量"], wb["充放电量"]
        assert (plan.max_row, plan.max_column) == (145, 2)
        assert (store.max_row, store.max_column) == (7, 5)
        assert plan["A1"].value == "时间段" and plan["B1"].value == "购电量"
        assert plan["A2"].value == "00:00-00:10" and plan["A145"].value == "23:50-24:00"
        assert plan["A62"].value == "10:00-10:10"
        for r in rows:
            cell = plan.cell(r["slot"] + 1, 2)
            assert cell.value == pytest.approx(r["grid_kwh"], abs=1e-12)
            assert cell.number_format == "0.00"
        template = load_workbook(TEMPLATE, read_only=True)
        try:
            tstore = template["充放电量"]
            for i in range(1, 8):
                assert store.cell(i, 1).value == tstore.cell(i, 1).value
                assert store.cell(i, 4).value == tstore.cell(i, 4).value
        finally:
            template.close()
        for b in range(6):
            part = rows[24 * b: 24 * (b + 1)]
            assert store.cell(b + 2, 2).value == pytest.approx(sum(r["charge_kwh"] for r in part), abs=1e-9)
            assert store.cell(b + 2, 3).value == pytest.approx(sum(r["discharge_kwh"] for r in part), abs=1e-9)
        assert store["E2"].value == pytest.approx(rows[0]["soc_start_kwh"])
        assert store["E3"].value == pytest.approx(rows[-1]["soc_end_kwh"])
        assert all(store.cell(i, j).value is None for i in range(4, 8) for j in (4, 5))
    finally:
        wb.close()


def test_template_mapping(run_dir):
    out, _ = run_dir
    with (out / "template_mapping.csv").open(encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream)
        assert reader.fieldnames == MAPPING_COLUMNS
        rows = list(reader)
    assert len(rows) == 144
    assert rows[0] == {"slot": "1", "input_excel_row": "2", "template_label_cell": "A2",
                       "original_label": "0:10-0:20", "corrected_label": "00:00-00:10"}
    assert rows[60]["template_label_cell"] == "A62"
    assert rows[60]["original_label"] == "10:10-10:20"
    assert rows[60]["corrected_label"] == "10:00-10:10"
    assert rows[-1]["original_label"] == "0:00+1-0:10+1"
    assert rows[-1]["corrected_label"] == "23:50-24:00"
    template = load_workbook(TEMPLATE, read_only=True)
    try:
        sheet = template["计划购电量"]
        for r in rows:
            assert sheet.cell(int(r["slot"]) + 1, 1).value == r["original_label"]
    finally:
        template.close()


def test_tables_md_matches_ledger(run_dir):
    out, summary = run_dir
    text = (out / "tables_q1.md").read_text(encoding="utf-8")
    rows = read_schedule_csv(out / "schedule.csv")
    for slot in [61, 73, 85, 97, 109, 121]:
        r = rows[slot - 1]
        assert f"| {r['start']}-{r['end']} | {r['grid_kwh']:.2f} |" in text
    assert f"{summary['objective_yuan']:.2f}" in text
    assert f"{summary['total_grid_kwh']:.2f}" in text


def test_cli_rejects_bad_config(tmp_path):
    config = json.loads((ROOT / "configs/q1.json").read_text(encoding="utf-8"))
    config["settlement_side"] = "battery"
    path = tmp_path / "bad.json"
    path.write_text(json.dumps(config), encoding="utf-8")
    assert q1.main(["--config", str(path), "--output", str(tmp_path / "out")]) == 1


def test_cli_reports_infeasible_solve(tmp_path):
    config = json.loads((ROOT / "configs/q1.json").read_text(encoding="utf-8"))
    config["initial_soc_kwh"] = 20000
    config["input"] = str((ROOT / config["input"]).resolve())
    config["template"] = str((ROOT / config["template"]).resolve())
    path = tmp_path / "infeasible.json"
    path.write_text(json.dumps(config), encoding="utf-8")
    assert q1.main(["--config", str(path), "--output", str(tmp_path / "out")]) == 1
    assert not (tmp_path / "out" / "schedule.csv").exists()

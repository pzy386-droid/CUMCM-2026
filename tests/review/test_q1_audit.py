"""Mutation tests for high-impact physical/accounting errors; no production imports."""
import copy
import datetime as dt
import subprocess
import sys

import numpy as np
import pytest

from review.q1_audit import (
    AuditError, DELTA, M, N, ROOT, check_optimal_cost, check_records,
    endpoint_minutes, make_records, read_input, solve_reference,
)


@pytest.fixture(scope="module")
def values():
    return read_input()


@pytest.fixture
def baseline(values):
    net = (values[:, 1]-values[:, 2])*DELTA
    return make_records(values, np.maximum(net, 0), np.zeros(N), np.zeros(N),
                        np.maximum(-net, 0), np.full(N, 6000), np.full(N, 6000))


def test_no_storage_is_a_feasible_control(values, baseline):
    metrics = check_records(baseline, values)
    assert metrics["objective_yuan"] > 0
    assert metrics["soc_initial_kwh"] == metrics["soc_final_kwh"] == 6000


@pytest.mark.parametrize("value,expected", [
    (dt.time(0, 10), 10), ("10:10", 610), ("0:00+1", 1440), ("24:00", 1440),
])
def test_mixed_timestamp_types(value, expected):
    assert endpoint_minutes(value) == expected


@pytest.mark.parametrize("corruption,error", [
    ("units", "energy balance"),
    ("time", "time start"),
    ("nonfinite", "non-finite"),
    ("cost", "per-slot cost"),
    ("initial", "initial/terminal"),
    ("power", "power bound"),
    ("efficiency", "SOC dynamics"),
    ("simultaneous", "simultaneous charge"),
])
def test_rejects_corruptions(values, baseline, corruption, error):
    bad = copy.deepcopy(baseline)
    if corruption == "units":
        bad[0]["grid_kwh"] *= 1.1
        bad[0]["cost_yuan"] = bad[0]["grid_kwh"]*bad[0]["price_yuan_per_kwh"]
    elif corruption == "time":
        bad[0]["start"] = "00:10"
    elif corruption == "nonfinite":
        bad[0]["grid_kwh"] = float("nan")
    elif corruption == "cost":
        bad[0]["cost_yuan"] += 1
    elif corruption == "initial":
        for row in bad:
            row["soc_start_kwh"] += 100
            row["soc_end_kwh"] += 100
    elif corruption == "power":
        bad[0]["charge_kwh"] = M+1
    elif corruption == "efficiency":
        bad[0]["soc_end_kwh"] += 1
        bad[1]["soc_start_kwh"] += 1
    elif corruption == "simultaneous":
        bad[0]["charge_kwh"] = 1
        bad[0]["discharge_kwh"] = .81
        bad[0]["grid_kwh"] += .19
        bad[0]["cost_yuan"] = bad[0]["grid_kwh"]*bad[0]["price_yuan_per_kwh"]
    with pytest.raises(AuditError, match=error):
        check_records(bad, values)


def test_lp_certificate_and_suboptimal_baseline(values, baseline):
    ref, rows = solve_reference(values)
    assert abs(ref["primal_dual_gap_yuan"]) < 1e-4
    assert ref["objective_yuan"] < check_records(baseline, values)["objective_yuan"]
    check_optimal_cost(ref["objective_yuan"], ref["objective_yuan"])
    with pytest.raises(AuditError, match="above LP bound"):
        check_optimal_cost(ref["objective_yuan"]+1, ref["objective_yuan"])
    with pytest.raises(AuditError, match="below certified"):
        check_optimal_cost(ref["objective_yuan"]-1, ref["objective_yuan"])


def test_missing_run_fails_closed(tmp_path):
    result = subprocess.run([sys.executable, str(ROOT/"review/verify_q1.py"), "--run", str(tmp_path)],
                            capture_output=True, text=True)
    assert result.returncode == 1
    assert '"status": "failed"' in result.stdout
    assert "missing required output" in result.stdout

def test_summary_and_nonfinite_gap(tmp_path, values, baseline):
    import json
    from review.q1_audit import check_summary, verify_manifest
    hashes = verify_manifest()
    metrics = check_records(baseline, values)
    keys = ["objective_yuan", "total_grid_kwh", "total_charge_kwh",
            "total_discharge_kwh", "total_spill_kwh", "soc_initial_kwh", "soc_final_kwh"]
    summary = {
        "schema_version": 1,
        **{key: metrics[key] for key in keys},
        "input_sha256": hashes["data/raw/附件1.xlsx"],
        "template_sha256": hashes["data/templates/result1.xlsx"],
        "solver": {"name": "unit-test-fixture", "status": "optimal", "mip_gap": 0, "wall_seconds": 0},
        "source_git_commit": "1"*40,
        "source_dirty": False,
    }
    path = tmp_path/"summary.json"
    path.write_text(json.dumps(summary), encoding="utf-8")
    check_summary(path, metrics, hashes)
    summary["solver"]["mip_gap"] = float("nan")
    path.write_text(json.dumps(summary), encoding="utf-8")
    with pytest.raises(AuditError, match="non-finite"):
        check_summary(path, metrics, hashes)


def test_original_blank_template_is_not_accepted(baseline):
    from review.q1_audit import check_workbook
    with pytest.raises(AuditError, match="Excel time label"):
        check_workbook(ROOT/"data/templates/result1.xlsx", baseline)


def test_template_mapping_and_corruption(tmp_path, baseline):
    import csv
    from openpyxl import load_workbook
    from review.q1_audit import check_mapping
    template = load_workbook(ROOT/"data/templates/result1.xlsx", read_only=True)
    try:
        original = [template["计划购电量"].cell(i+2, 1).value for i in range(N)]
    finally:
        template.close()
    fields = ["slot", "input_excel_row", "template_label_cell", "original_label", "corrected_label"]
    rows = [
        dict(zip(fields, [i+1, i+2, f"A{i+2}", original[i], f"{row['start']}-{row['end']}"]))
        for i, row in enumerate(baseline)
    ]
    path = tmp_path/"mapping.csv"
    def save():
        with path.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)
    save()
    check_mapping(path, baseline)
    rows[0]["input_excel_row"] = 3
    save()
    with pytest.raises(AuditError, match="mapping indices"):
        check_mapping(path, baseline)

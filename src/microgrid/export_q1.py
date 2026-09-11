"""Ledger construction and export of the Q1 deliverables (docs/q1_spec.md §输出合同).

Every artefact is generated from the single per-slot ledger returned by
``build_ledger``; nothing is re-solved or hand-edited downstream.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .inputs import Q1Input, sha256_file
from .model_q1 import Q1Params, Q1Solution

LEDGER_COLUMNS = [
    "slot", "start", "end", "price_yuan_per_kwh", "load_kw", "pv_kw",
    "grid_kwh", "charge_kwh", "discharge_kwh", "spill_kwh",
    "soc_start_kwh", "soc_end_kwh", "cost_yuan",
]
MAPPING_COLUMNS = ["slot", "input_excel_row", "template_label_cell", "original_label", "corrected_label"]
SELECTED_SLOTS = [61, 73, 85, 97, 109, 121]   # 10:00, 12:00, 14:00, 16:00, 18:00, 20:00
PLAN_SHEET = "计划购电量"
STORAGE_SHEET = "充放电量"
DISPLAY_FORMAT = "0.00"


class ExportError(RuntimeError):
    pass


def build_ledger(data: Q1Input, sol: Q1Solution) -> list[dict[str, Any]]:
    rows = []
    for i in range(data.slots):
        slot = i + 1
        rows.append({
            "slot": slot,
            "start": data.start_label(slot),
            "end": data.end_label(slot),
            "price_yuan_per_kwh": float(data.price[i]),
            "load_kw": float(data.load_kw[i]),
            "pv_kw": float(data.pv_kw[i]),
            "grid_kwh": float(sol.grid_kwh[i]),
            "charge_kwh": float(sol.charge_kwh[i]),
            "discharge_kwh": float(sol.discharge_kwh[i]),
            "spill_kwh": float(sol.spill_kwh[i]),
            "soc_start_kwh": float(sol.soc_kwh[i]),
            "soc_end_kwh": float(sol.soc_kwh[i + 1]),
            "cost_yuan": float(data.price[i] * sol.grid_kwh[i]),
        })
    return rows


def ledger_totals(rows: list[dict[str, Any]]) -> dict[str, float]:
    return {
        "objective_yuan": float(sum(r["cost_yuan"] for r in rows)),
        "total_grid_kwh": float(sum(r["grid_kwh"] for r in rows)),
        "total_charge_kwh": float(sum(r["charge_kwh"] for r in rows)),
        "total_discharge_kwh": float(sum(r["discharge_kwh"] for r in rows)),
        "total_spill_kwh": float(sum(r["spill_kwh"] for r in rows)),
        "soc_initial_kwh": float(rows[0]["soc_start_kwh"]),
        "soc_final_kwh": float(rows[-1]["soc_end_kwh"]),
    }


def four_hour_blocks(rows: list[dict[str, Any]], slots_per_block: int = 24) -> list[dict[str, Any]]:
    blocks = []
    for b in range(0, len(rows), slots_per_block):
        part = rows[b:b + slots_per_block]
        blocks.append({
            "interval": f"{part[0]['start']}-{part[-1]['end']}",
            "slots": f"{part[0]['slot']}-{part[-1]['slot']}",
            "charge_kwh": float(sum(r["charge_kwh"] for r in part)),
            "discharge_kwh": float(sum(r["discharge_kwh"] for r in part)),
        })
    return blocks


def selected_intervals(rows: list[dict[str, Any]], slots=SELECTED_SLOTS) -> list[dict[str, Any]]:
    out = []
    for s in slots:
        r = rows[s - 1]
        out.append({"slot": s, "interval": f"{r['start']}-{r['end']}",
                    "input_excel_row": s + 1, "grid_kwh": r["grid_kwh"], "cost_yuan": r["cost_yuan"]})
    return out


def write_schedule_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    with Path(path).open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=LEDGER_COLUMNS)
        writer.writeheader()
        for r in rows:
            writer.writerow({k: (repr(v) if isinstance(v, float) else v) for k, v in r.items()})


def read_schedule_csv(path: Path) -> list[dict[str, Any]]:
    with Path(path).open(encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream)
        if reader.fieldnames != LEDGER_COLUMNS:
            raise ExportError("schedule.csv columns differ from contract")
        rows = []
        for r in reader:
            row = dict(r)
            row["slot"] = int(row["slot"])
            for k in LEDGER_COLUMNS[3:]:
                row[k] = float(row[k])
            rows.append(row)
    return rows


def template_mapping(template_path: Path, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Original label of each plan-sheet cell A2:A145 versus the corrected day interval."""
    wb = load_workbook(template_path, read_only=True, data_only=False)
    try:
        sheet = wb[PLAN_SHEET]
        originals = [sheet.cell(i + 2, 1).value for i in range(len(rows))]
    finally:
        wb.close()
    mapping = []
    for i, r in enumerate(rows):
        mapping.append({
            "slot": r["slot"], "input_excel_row": r["slot"] + 1,
            "template_label_cell": f"A{r['slot'] + 1}",
            "original_label": originals[i],
            "corrected_label": f"{r['start']}-{r['end']}",
        })
    return mapping


def write_template_mapping_csv(path: Path, mapping: list[dict[str, Any]]) -> None:
    with Path(path).open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=MAPPING_COLUMNS)
        writer.writeheader()
        writer.writerows(mapping)


def write_result1_xlsx(template_path: Path, out_path: Path, rows: list[dict[str, Any]],
                       params: Q1Params) -> None:
    """Fill a copy of the template: corrected A2:A145 labels, g in B, block sums, S0/S_N."""
    template_path, out_path = Path(template_path), Path(out_path)
    n = len(rows)
    wb = load_workbook(template_path, data_only=False)
    try:
        if wb.sheetnames != [PLAN_SHEET, STORAGE_SHEET]:
            raise ExportError(f"template sheets {wb.sheetnames} differ from contract")
        plan = wb[PLAN_SHEET]
        if (plan.max_row, plan.max_column) != (n + 1, 2):
            raise ExportError("template plan sheet dimensions differ from contract")
        for r in rows:
            excel_row = r["slot"] + 1
            plan.cell(excel_row, 1).value = f"{r['start']}-{r['end']}"
            cell = plan.cell(excel_row, 2)
            cell.value = r["grid_kwh"]
            cell.number_format = DISPLAY_FORMAT
        store = wb[STORAGE_SHEET]
        if (store.max_row, store.max_column) != (7, 5):
            raise ExportError("template storage sheet dimensions differ from contract")
        blocks = four_hour_blocks(rows)
        if len(blocks) != 6:
            raise ExportError("expected six four-hour blocks")
        for b, block in enumerate(blocks):
            excel_row = b + 2
            for col, key in ((2, "charge_kwh"), (3, "discharge_kwh")):
                cell = store.cell(excel_row, col)
                cell.value = block[key]
                cell.number_format = DISPLAY_FORMAT
        for excel_row, value in ((2, rows[0]["soc_start_kwh"]), (3, rows[-1]["soc_end_kwh"])):
            cell = store.cell(excel_row, 5)
            cell.value = value
            cell.number_format = DISPLAY_FORMAT
        out_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(out_path)
    finally:
        wb.close()
    _readback_check(template_path, out_path, rows)


def _readback_check(template_path: Path, out_path: Path, rows: list[dict[str, Any]], tol=1e-9) -> None:
    wb = load_workbook(out_path, data_only=False)
    tpl = load_workbook(template_path, read_only=True, data_only=False)
    try:
        if wb.sheetnames != [PLAN_SHEET, STORAGE_SHEET]:
            raise ExportError("exported sheets changed")
        plan, store = wb[PLAN_SHEET], wb[STORAGE_SHEET]
        if (plan.max_row, plan.max_column) != (len(rows) + 1, 2) or (store.max_row, store.max_column) != (7, 5):
            raise ExportError("exported dimensions changed")
        for r in rows:
            i = r["slot"] + 1
            if plan.cell(i, 1).value != f"{r['start']}-{r['end']}" or abs(plan.cell(i, 2).value - r["grid_kwh"]) > tol:
                raise ExportError(f"plan sheet readback mismatch at row {i}")
        tstore = tpl[STORAGE_SHEET]
        for i in range(1, 8):
            if store.cell(i, 1).value != tstore.cell(i, 1).value or store.cell(i, 4).value != tstore.cell(i, 4).value:
                raise ExportError("storage sheet labels changed")
        for b, block in enumerate(four_hour_blocks(rows)):
            if abs(store.cell(b + 2, 2).value - block["charge_kwh"]) > tol or \
               abs(store.cell(b + 2, 3).value - block["discharge_kwh"]) > tol:
                raise ExportError("storage sheet readback mismatch")
        if abs(store["E2"].value - rows[0]["soc_start_kwh"]) > tol or abs(store["E3"].value - rows[-1]["soc_end_kwh"]) > tol:
            raise ExportError("SOC cells readback mismatch")
        for i in range(4, 8):
            if store.cell(i, 5).value is not None:
                raise ExportError("storage sheet blank cells were filled")
    finally:
        wb.close()
        tpl.close()


def _fmt(value: float) -> str:
    return f"{value:.2f}"


def write_tables_md(path: Path, rows: list[dict[str, Any]], totals: dict[str, float]) -> None:
    sel = selected_intervals(rows)
    blocks = four_hour_blocks(rows)
    lines = [
        "# 问题1 论文表格（由 schedule.csv 同一账本生成，显示两位小数，底层全精度见 schedule.csv / summary.json）",
        "",
        "## 表1 微网在指定时间段的购电量及全天的购电量和购电费",
        "",
        "| 时间段 | 购电量（kWh） | 时段序号 t | 附件1输入行 |",
        "|---|---:|---:|---:|",
    ]
    for s in sel:
        lines.append(f"| {s['interval']} | {_fmt(s['grid_kwh'])} | {s['slot']} | {s['input_excel_row']} |")
    lines += [
        "",
        f"全天购电量：{_fmt(totals['total_grid_kwh'])} kWh；全天购电费：{_fmt(totals['objective_yuan'])} 元。",
        "",
        "## 表2 储能设备在指定时间段的充放电量及 0:00 和 24:00 的储电量",
        "",
        "| 时间段 | 充电量（kWh） | 放电量（kWh） |",
        "|---|---:|---:|",
    ]
    for b in blocks:
        lines.append(f"| {b['interval']} | {_fmt(b['charge_kwh'])} | {_fmt(b['discharge_kwh'])} |")
    lines += [
        "",
        f"0:00 储电量：{_fmt(totals['soc_initial_kwh'])} kWh；24:00 储电量：{_fmt(totals['soc_final_kwh'])} kWh。",
        "",
        "说明：充电量/放电量为母线侧电量；储电量按 S_t = S_(t-1) + 0.9c_t − d_t/0.9 更新；"
        f"全天充电 {_fmt(totals['total_charge_kwh'])} kWh、放电 {_fmt(totals['total_discharge_kwh'])} kWh、"
        f"富余未利用 {_fmt(totals['total_spill_kwh'])} kWh。",
        "",
    ]
    Path(path).write_text("\n".join(lines), encoding="utf-8")


def write_json(path: Path, payload: dict[str, Any]) -> None:
    Path(path).write_text(json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
                          encoding="utf-8")


def copy_input_hash(path: Path) -> str:
    return sha256_file(path)

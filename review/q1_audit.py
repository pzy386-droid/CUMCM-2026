"""Independent Q1 audit. This module does not import microgrid production code."""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import json
import math
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import lil_matrix

N = 144
DELTA = 1 / 6
ETA = 0.9
M = 5000 / 6
TOL = 1e-6
ROOT = Path(__file__).resolve().parents[1]
COLUMNS = [
    "slot", "start", "end", "price_yuan_per_kwh", "load_kw", "pv_kw",
    "grid_kwh", "charge_kwh", "discharge_kwh", "spill_kwh",
    "soc_start_kwh", "soc_end_kwh", "cost_yuan",
]


class AuditError(ValueError):
    pass


def require(condition, message):
    if not condition:
        raise AuditError(message)


def scalar(value, label):
    require(not isinstance(value, bool), f"{label}: boolean is not a number")
    try:
        result = float(value)
    except (TypeError, ValueError):
        raise AuditError(f"{label}: expected number") from None
    require(math.isfinite(result), f"{label}: non-finite number")
    return result


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def time_label(minutes):
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def endpoint_minutes(value):
    if isinstance(value, dt.time):
        require(value.second == 0 and value.microsecond == 0, "time includes seconds")
        return value.hour * 60 + value.minute
    if isinstance(value, str):
        text = value.strip()
        if text == "0:00+1":
            return 1440
        pieces = text.split(":")
        require(len(pieces) == 2, f"invalid timestamp: {value}")
        hour, minute = map(int, pieces)
        require(0 <= hour <= 24 and 0 <= minute < 60, "invalid hour/minute")
        require(hour != 24 or minute == 0, "timestamp past 24:00")
        return hour * 60 + minute
    raise AuditError(f"unsupported timestamp type: {type(value).__name__}")


def verify_manifest(root=ROOT):
    root = Path(root)
    doc = json.loads((root / "data/manifest.json").read_text(encoding="utf-8"))
    require(doc.get("schema_version") == 1, "manifest schema")
    hashes = {}
    for entry in doc["files"]:
        rel = entry["path"]
        target = (root / rel).resolve()
        require(target.is_relative_to(root.resolve()), "manifest path escapes repository")
        require(target.is_file(), f"missing source: {rel}")
        require(sha256(target) == entry["sha256"], f"source hash changed: {rel}")
        hashes[rel] = entry["sha256"]
    return hashes


def read_input(root=ROOT):
    root = Path(root)
    verify_manifest(root)
    wb = load_workbook(root / "data/raw/附件1.xlsx", read_only=True, data_only=False)
    try:
        require(len(wb.worksheets) == 1, "input sheet count")
        rows = list(wb.worksheets[0].values)
    finally:
        wb.close()
    require(len(rows) == 145 and all(len(row) == 4 for row in rows), "input dimensions")
    require(list(rows[0]) == ["时间", "电价", "小区负载", "光伏发电预测功率"], "input headers")
    for i, row in enumerate(rows[1:], 1):
        require(endpoint_minutes(row[0]) == 10 * i, f"input time at slot {i}")
    values = np.array([[scalar(v, "input") for v in row[1:]] for row in rows[1:]])
    require(np.all(values[:, 0] > 0), "Q1 requires positive input prices")
    require(np.all(values[:, 1:] >= 0), "negative input power")
    return values  # price, load kW, PV kW


def solve_reference(values):
    """A lower bound with no mode constraints, assembled independently."""
    total = 5 * N + 1
    g, c, d, r, state = 0, N, 2 * N, 3 * N, 4 * N
    matrix = lil_matrix((2 * N, total), dtype=float)
    rhs = np.zeros(2 * N)
    for i in range(N):
        matrix[i, g+i] = 1
        matrix[i, c+i] = -1
        matrix[i, d+i] = 1
        matrix[i, r+i] = -1
        rhs[i] = (values[i, 1] - values[i, 2]) * DELTA
        matrix[N+i, state+i+1] = 1
        matrix[N+i, state+i] = -1
        matrix[N+i, c+i] = -ETA
        matrix[N+i, d+i] = 1 / ETA
    bounds = ([(0, None)] * N + [(0, M)] * (2*N)
              + [(0, None)] * N + [(1200, 10800)] * (N+1))
    bounds[state] = (6000, 6000)
    bounds[state+N] = (6000, 6000)
    objective = np.zeros(total)
    objective[:N] = values[:, 0]
    matrix = matrix.tocsr()
    result = linprog(objective, A_eq=matrix, b_eq=rhs, bounds=bounds, method="highs")
    require(result.success and result.status == 0, f"reference LP failed: {result.message}")
    dual = float(rhs @ result.eqlin.marginals)
    for i, (lower, upper) in enumerate(bounds):
        if lower is not None:
            dual += lower * result.lower.marginals[i]
        if upper is not None:
            dual += upper * result.upper.marginals[i]
    stationarity = (objective - matrix.T @ result.eqlin.marginals
                    - result.lower.marginals - result.upper.marginals)
    require(abs(result.fun - dual) <= cost_tolerance(result.fun), "LP primal/dual gap")
    require(np.max(np.abs(matrix @ result.x-rhs)) < TOL, "LP primal residual")
    require(np.max(np.abs(stationarity)) < TOL, "LP dual stationarity")
    x = result.x
    records = make_records(values, x[:N], x[N:2*N], x[2*N:3*N],
                           x[3*N:4*N], x[4*N:4*N+N], x[4*N+1:4*N+N+1])
    no_storage = float(values[:, 0] @ np.maximum((values[:, 1]-values[:, 2])*DELTA, 0))
    certificate = {
        "purpose": "independent_reference_not_official_result1",
        "lp_status": "optimal",
        "objective_yuan": float(result.fun),
        "dual_bound_yuan": dual,
        "primal_dual_gap_yuan": float(result.fun-dual),
        "max_primal_residual": float(np.max(np.abs(matrix @ result.x-rhs))),
        "max_dual_stationarity_residual": float(np.max(np.abs(stationarity))),
        "no_storage_cost_yuan": no_storage,
    }
    try:
        certificate["feasible_under_full_q1_rules"] = True
        certificate["schedule_metrics"] = check_records(records, values)
    except AuditError as exc:
        certificate["feasible_under_full_q1_rules"] = False
        certificate["feasibility_note"] = str(exc)
    return certificate, records


def make_records(values, grid, charge, discharge, spill, start_soc, end_soc):
    records = []
    for i in range(N):
        records.append(dict(zip(COLUMNS, [
            i+1, time_label(i*10), time_label((i+1)*10),
            *map(float, values[i]), float(grid[i]), float(charge[i]),
            float(discharge[i]), float(spill[i]), float(start_soc[i]), float(end_soc[i]),
            float(values[i, 0] * grid[i]),
        ])))
    return records


def write_records(path, records):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(records)


def read_records(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        require(reader.fieldnames == COLUMNS, "schedule columns/order")
        records = list(reader)
    return records


def check_records(records, values):
    require(len(records) == N, "schedule must have 144 rows")
    numeric = COLUMNS[3:]
    a = np.zeros((N, len(numeric)))
    for i, row in enumerate(records):
        require(scalar(row["slot"], "slot") == i+1, f"slot order at row {i+1}")
        require(row["start"] == time_label(i*10), f"time start at slot {i+1}")
        require(row["end"] == time_label((i+1)*10), f"time end at slot {i+1}")
        a[i] = [scalar(row[k], f"slot {i+1} {k}") for k in numeric]
    require(np.max(np.abs(a[:, :3]-values)) <= 1e-9, "input columns differ from attachment")
    price, load, pv, grid, charge, discharge, spill, s0, s1, cost = a.T
    require(np.min(a[:, 3:7]) >= -TOL, "negative energy")
    require(max(charge.max(), discharge.max()) <= M+TOL, "charge/discharge power bound")
    require(np.all(grid <= load*DELTA+M+TOL), "grid upper bound")
    require(np.all((charge <= TOL) | (discharge <= TOL)), "simultaneous charge/discharge")
    require(np.all((discharge <= TOL) | (spill <= TOL)), "discharge and spill simultaneously")
    require(min(s0.min(), s1.min()) >= 1200-TOL, "SOC lower bound")
    require(max(s0.max(), s1.max()) <= 10800+TOL, "SOC upper bound")
    require(abs(s0[0]-6000) <= TOL and abs(s1[-1]-6000) <= TOL, "initial/terminal SOC")
    require(np.max(np.abs(s0[1:]-s1[:-1])) <= TOL, "SOC continuity")
    energy = grid + pv*DELTA + discharge-load*DELTA-charge-spill
    dynamics = s1-s0-ETA*charge+discharge/ETA
    require(np.max(np.abs(energy)) <= TOL, "energy balance / kW-kWh conversion")
    require(np.max(np.abs(dynamics)) <= TOL, "SOC dynamics / efficiency")
    require(np.max(np.abs(cost-price*grid)) <= 1e-6, "per-slot cost")
    return {
        "objective_yuan": float(price @ grid),
        "total_grid_kwh": float(grid.sum()),
        "total_charge_kwh": float(charge.sum()),
        "total_discharge_kwh": float(discharge.sum()),
        "total_spill_kwh": float(spill.sum()),
        "soc_initial_kwh": float(s0[0]),
        "soc_final_kwh": float(s1[-1]),
        "max_energy_residual_kwh": float(np.max(np.abs(energy))),
        "max_soc_residual_kwh": float(np.max(np.abs(dynamics))),
    }


def cost_tolerance(cost):
    return max(1e-5, 1e-8*abs(float(cost)))


def check_optimal_cost(cost, reference):
    tol = cost_tolerance(reference)
    require(cost >= reference-tol, "cost below certified LP bound: inconsistent model/data")
    require(cost <= reference+tol, "cost above LP bound: optimality needs further review")


def check_summary(path, metrics, hashes):
    summary = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    require(summary.get("schema_version") == 1, "summary schema")
    for key in ["objective_yuan", "total_grid_kwh", "total_charge_kwh",
                "total_discharge_kwh", "total_spill_kwh", "soc_initial_kwh", "soc_final_kwh"]:
        require(abs(scalar(summary.get(key), key)-metrics[key]) <= 1e-6, f"summary {key}")
    require(summary.get("input_sha256") == hashes["data/raw/附件1.xlsx"], "summary input hash")
    require(summary.get("template_sha256") == hashes["data/templates/result1.xlsx"], "summary template hash")
    solver = summary.get("solver", {})
    require(solver.get("status") == "optimal", "solver must report actual optimal status")
    gap = scalar(solver.get("mip_gap"), "mip_gap")
    require(0 <= gap <= 1e-8, "solver MIP gap")
    require(isinstance(solver.get("name"), str) and solver["name"], "solver name")
    require(scalar(solver.get("wall_seconds"), "wall_seconds") >= 0, "negative wall time")
    commit = summary.get("source_git_commit", "")
    require(isinstance(commit, str) and len(commit) == 40 and
            all(ch in "0123456789abcdef" for ch in commit.lower()), "source git commit")
    require(isinstance(summary.get("source_dirty"), bool), "source_dirty must be boolean")
    return summary


def near_cell(cell, expected, label):
    require(abs(scalar(cell.value, label)-float(expected)) <= TOL, f"Excel {label}")


def check_workbook(path, records, root=ROOT):
    wb = load_workbook(path, data_only=False)
    template = load_workbook(Path(root) / "data/templates/result1.xlsx", data_only=False)
    try:
        require(wb.sheetnames == ["计划购电量", "充放电量"], "Excel sheet names/order")
        plan = wb["计划购电量"]
        require((plan.max_row, plan.max_column) == (145, 2), "Excel plan dimensions")
        require([plan.cell(1, j).value for j in (1, 2)] == ["时间段", "购电量"], "Excel plan headers")
        for i, row in enumerate(records, 2):
            require(plan.cell(i, 1).value == f"{row['start']}-{row['end']}", f"Excel time label A{i}")
            near_cell(plan.cell(i, 2), row["grid_kwh"], f"plan B{i}")
        store = wb["充放电量"]
        original = template["充放电量"]
        require((store.max_row, store.max_column) == (7, 5), "Excel storage dimensions")
        require([store.cell(1, j).value for j in range(1, 6)] ==
                [original.cell(1, j).value for j in range(1, 6)], "Excel storage headers")
        for block in range(6):
            i = block+2
            require(store.cell(i, 1).value == original.cell(i, 1).value, "Excel four-hour labels")
            part = records[block*24:(block+1)*24]
            near_cell(store.cell(i, 2), sum(float(x["charge_kwh"]) for x in part), f"storage B{i}")
            near_cell(store.cell(i, 3), sum(float(x["discharge_kwh"]) for x in part), f"storage C{i}")
        require(endpoint_minutes(store["D2"].value) == 0, "Excel initial time")
        require(endpoint_minutes(store["D3"].value) == 1440, "Excel final time")
        near_cell(store["E2"], records[0]["soc_start_kwh"], "initial SOC E2")
        near_cell(store["E3"], records[-1]["soc_end_kwh"], "final SOC E3")
        require(all(store.cell(i, j).value is None for i in range(4, 8) for j in (4, 5)),
                "Excel template blanks changed")
    finally:
        wb.close()
        template.close()


def check_mapping(path, records, root=ROOT):
    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        require(reader.fieldnames == ["slot", "input_excel_row", "template_label_cell",
                                      "original_label", "corrected_label"], "mapping columns")
        rows = list(reader)
    require(len(rows) == N, "mapping row count")
    wb = load_workbook(Path(root) / "data/templates/result1.xlsx", read_only=True, data_only=False)
    try:
        sheet = wb["计划购电量"]
        for i, (mapping, record) in enumerate(zip(rows, records), 1):
            require(int(mapping["slot"]) == i and int(mapping["input_excel_row"]) == i+1, "mapping indices")
            require(mapping["template_label_cell"] == f"A{i+1}", "mapping template cell")
            require(mapping["original_label"] == sheet.cell(i+1, 1).value, "mapping original label")
            require(mapping["corrected_label"] == f"{record['start']}-{record['end']}", "mapping corrected label")
    finally:
        wb.close()

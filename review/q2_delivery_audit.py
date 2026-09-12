"""Read-only independent checks of the small Q2 delivery, not full-run acceptance."""
import argparse
import datetime as dt
import hashlib
import json
import math
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--delivery', type=Path, required=True)
    p.add_argument('--repo', type=Path, required=True)
    p.add_argument('--report', type=Path, required=True)
    a = p.parse_args()
    summary = json.loads((a.delivery / 'summary.json').read_text(encoding='utf-8'))
    manifest = json.loads((a.delivery / 'ARCHIVE_MANIFEST_full.json').read_text(encoding='utf-8'))
    sp = summary['submission_period']
    checks, findings = [], []

    def check(name, condition, detail=None):
        checks.append(dict(check=name, passed=bool(condition), detail=detail))

    expected_files = manifest['runs']['main-no-intraday-year']['files']
    for f in ['result2.xlsx', 'summary.json', 'tables_q2.md']:
        digest = hashlib.sha256((a.delivery / f).read_bytes()).hexdigest()
        check('manifest ' + f, digest == expected_files[f]['sha256'], digest)
    price_wb = load_workbook(a.repo / 'data/raw/附件1.xlsx', read_only=True, data_only=True)
    prices = [r[1] for r in list(price_wb.active.values)[1:]]
    price_wb.close()
    actual_wb = load_workbook(a.repo / 'data/raw/附件2.xlsx', read_only=True, data_only=True)
    actual = {}
    for sheet in actual_wb:
        actual[sheet.title] = {r[0].date(): sum(r[1:]) / 6 for r in list(sheet.values)[1:]}
    actual_wb.close()
    wb = load_workbook(a.delivery / 'result2.xlsx', data_only=False)
    check('sheet names', wb.sheetnames == ['计划购电量', '充放电量', '紧急购电量'])
    for sheet in wb:
        bad = [c.coordinate for row in sheet for c in row if c.data_type in ('f', 'e') or
               (isinstance(c.value, (int, float)) and not math.isfinite(c.value))]
        check('finite literal values ' + sheet.title, not bad, bad[:10])
    ws = wb['计划购电量']
    check('plan dimensions', (ws.max_row, ws.max_column) == (335, 147))
    def hm(m):
        return f'{m//60:02d}:{m%60:02d}'
    check('all 144 labels', all(ws.cell(1, k+2).value == hm(k*10)+'-'+hm((k+1)*10) for k in range(144)))
    dates = [dt.date(2025, 2, 1) + dt.timedelta(days=i) for i in range(334)]
    day = {}
    negatives = []
    for i, date in enumerate(dates, 2):
        vals = [ws.cell(i, k+2).value for k in range(144)]
        check('date ' + str(date), ws.cell(i, 1).value.date() == date)
        for k, v in enumerate(vals):
            if v < 0:
                negatives.append(dict(date=str(date), slot=k+1, cell=ws.cell(i,k+2).coordinate, value=v))
        check('daily plan sum '+str(date), abs(sum(vals)-ws.cell(i,146).value) < 1e-6)
        day[date] = dict(g=sum(vals), gcost=sum(x*y for x,y in zip(vals,prices)), cost=ws.cell(i,147).value, values=vals)
    for key, reported in [('g','planned_grid_kwh'),('gcost','planned_cost_yuan'),('cost','total_cost_yuan')]:
        total = sum(d[key] for d in day.values())
        check('summary ' + reported, abs(total-sp[reported]) < 1e-5, total)
    check('nonnegative within numerical tolerance', all(x['value'] >= -1e-6 for x in negatives))
    if negatives:
        findings.append(dict(issue='negative contract quantities at numerical tolerance', count=len(negatives), minimum=min(x['value'] for x in negatives), examples=negatives[:10]))
    ws = wb['紧急购电量']
    ev = defaultdict(list)
    date = None
    for row in list(ws.values)[1:]:
        if row[0] is not None:
            date = row[0].date()
        if not row[1]:
            check('empty emergency row', row[2] == 0)
            continue
        m = re.fullmatch(r'(\d{1,2}):(\d{2})-(\d{1,2}):(\d{2})', row[1])
        assert m, row
        h1,m1,h2,m2 = map(int,m.groups())
        lo,hi = h1*60+m1,h2*60+m2
        check('event bounds', date in day and 0<=lo<hi<=1440 and lo%10==hi%10==0 and row[2]>0)
        ev[date].append((lo//10,hi//10,row[2]))
    for date,d in day.items():
        entries=ev[date]
        check('ordered nonoverlapping separated events '+str(date), all(x[1]<y[0] for x,y in zip(entries,entries[1:])))
        d['e']=sum(x[2] for x in entries)
        low=sum(5*min(prices[lo:hi])*energy for lo,hi,energy in entries)
        high=sum(5*max(prices[lo:hi])*energy for lo,hi,energy in entries)
        implied=d['cost']-d['gcost']
        check('emergency fee necessary bounds '+str(date), low-1e-6<=implied<=high+1e-6)
    check('emergency total', abs(sum(d['e'] for d in day.values())-sp['emergency_grid_kwh'])<1e-5, sum(d['e'] for d in day.values()))
    check('emergency event count', sum(map(len,ev.values()))==sp['emergency_events'],sum(map(len,ev.values())))
    ws=wb['充放电量']
    check('storage dimensions', (ws.max_row,ws.max_column)==(2005,6))
    prev=None
    max_res=0
    for i,date in enumerate(dates):
        base=2+6*i
        start,end=ws.cell(base,6).value,ws.cell(base+1,6).value
        check('storage date '+str(date),ws.cell(base,1).value.date()==date)
        if prev is not None:
            check('cross-day displayed SOC '+str(date),abs(start-prev)<1e-6)
        prev=end
        c=dch=0
        S=start
        for b in range(6):
            ch,dis=ws.cell(base+b,3).value,ws.cell(base+b,4).value
            check('block energy bounds',0<=ch<=20000+1e-6 and 0<=dis<=20000+1e-6)
            c+=ch;dch+=dis;S+=.9*ch-dis/.9
            check('four-hour boundary SOC',1200-1e-6<=S<=10800+1e-6)
        max_res=max(max_res,abs(S-end))
        check('day SOC dynamics '+str(date),abs(S-end)<1e-6)
        d=day[date]
        d.update(c=c,d=dch,s0=start,s1=end)
        d['spill']=d['g']+d['e']+actual['光伏发电实际功率'][date]+dch-actual['小区负载'][date]-c
        check('daily implied surplus nonnegative '+str(date),d['spill']>=-1e-6)
    for key,reported in [('c','charge_kwh'),('d','discharge_kwh'),('spill','spill_kwh')]:
        check('summary '+reported,abs(sum(d[key] for d in day.values())-sp[reported])<1e-4,sum(d[key] for d in day.values()))
    check('initial submission SOC',abs(day[dates[0]]['s0']-sp['soc_start_kwh'])<1e-6)
    check('final submission SOC',abs(day[dates[-1]]['s1']-sp['soc_end_kwh'])<1e-6)
    wb.close()
    selected={}
    md=(a.delivery/'tables_q2.md').read_text(encoding='utf-8')
    for ds in ['2025-03-20','2025-06-21','2025-09-23','2025-12-21']:
        d=day[dt.date.fromisoformat(ds)]
        section=md.split('## '+ds)[1].split('\n## ')[0]
        totals=dict(planned_grid_kwh=d['g'],planned_cost_yuan=d['gcost'],total_cost_yuan=d['cost'],emergency_kwh=d['e'],soc_start=d['s0'],soc_end=d['s1'])
        check('selected day displayed totals '+ds,all(f'{v:.2f}' in section for v in totals.values()))
        selected[ds]=totals
    fork_path=a.repo/'reports/q2/fork_main-no-intraday-year.json'
    fork=json.loads(fork_path.read_text(encoding='utf-8'))
    fork_soc={}
    for k in ['36','72','108']:
        fork_soc[k]={v:sum(r[k][v]['soc_end']-r[k]['keep']['soc_end'] for r in fork['per_day'])/len(fork['per_day']) for v in ['resolve','reselect']}
    report=dict(scope='Independent ZIP/workbook/static-report audit only; full run archive absent. Fee bounds are necessary checks, not per-slot validation.',
                status='passed_with_findings' if all(x['passed'] for x in checks) else 'failed',
                checks_run=len(checks),failed_checks=[x for x in checks if not x['passed']],findings=findings,
                totals={key:sum(d[key] for d in day.values()) for key in ['g','gcost','cost','e','c','d','spill']},
                max_day_soc_residual=max_res,selected_days=selected,fork_mean_terminal_soc_difference=fork_soc,
                checks_passed=sum(x['passed'] for x in checks))
    a.report.parent.mkdir(parents=True,exist_ok=True)
    a.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k!='checks'},ensure_ascii=False,indent=2))
    return 0 if not report['failed_checks'] else 1


if __name__=='__main__':
    raise SystemExit(main())
